import subprocess
import urllib.request
import urllib.parse
import json
import os
import openpyxl
import datetime

# --- Configuration & Paths ---
spreadsheet_id = "1Tznbdor6-JFLkGNtuN5StasMSopnLkhdqzgbq7NWdAU"
rev_xlsx_path = r"%USERPROFILE%\.gemini\antigravity\brain\agy-session-2287\scratch\revision_115.xlsx"
analyzed_json_path = r"<repo>\Google Arts & Culture\favorites_analyzed.json"
enriched_json_path = r"<repo>\Google Arts & Culture\favorites_enriched.json"

def get_token():
    return subprocess.check_output('gcloud auth print-access-token', shell=True, text=True).strip()

token = get_token()

def api_get(url):
    req = urllib.request.Request(url, headers={'Authorization': f'Bearer {token}'})
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read().decode())

def api_post(url, body):
    data = json.dumps(body).encode("utf-8")
    req = urllib.request.Request(
        url,
        data=data,
        headers={'Authorization': f'Bearer {token}', 'Content-Type': 'application/json'},
        method="POST"
    )
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read().decode())

print("=== STEP 1: Inspect Current Spreadsheet State ===")
meta = api_get(f"https://sheets.googleapis.com/v4/spreadsheets/{spreadsheet_id}?fields=sheets(properties,bandedRanges)")
existing_sheets = {s['properties']['sheetId']: s['properties']['title'] for s in meta['sheets']}
print("Existing sheets:", existing_sheets)

# Identify old banding if present on sheet 0
stale_bandings = []
for s in meta['sheets']:
    if s['properties']['sheetId'] == 0:
        for b in s.get('bandedRanges', []):
            stale_bandings.append(b['bandedRangeId'])
print(f"Stale bandings on Sheet ID 0: {stale_bandings}")

# Check if Catalogue Precepts tab already exists with a different sheetId
precepts_sheet_id = None
for s in meta['sheets']:
    if s['properties']['title'] == "Catalogue Precepts" and s['properties']['sheetId'] != 0:
        precepts_sheet_id = s['properties']['sheetId']

if precepts_sheet_id is None:
    # Dedicated new sheet ID
    precepts_sheet_id = 98234710
print(f"Catalogue Precepts will use Sheet ID: {precepts_sheet_id}")

print("\n=== STEP 2: Restore Sheet ID 0 Title & Structure ===")
batch_reqs = []

# If Sheet ID 0 has stale bandings, delete them
for bid in stale_bandings:
    batch_reqs.append({"deleteBanding": {"bandedRangeId": bid}})

# Rename Sheet ID 0 to 'Enriched Favorites' and configure grid properties
batch_reqs.append({
    "updateSheetProperties": {
        "properties": {
            "sheetId": 0,
            "title": "Enriched Favorites",
            "gridProperties": {
                "rowCount": 822,
                "columnCount": 49,
                "frozenRowCount": 1
            }
        },
        "fields": "title,gridProperties"
    }
})

# Add brand new sheet for Catalogue Precepts if not already created
existing_sheet_ids = [s['properties']['sheetId'] for s in meta['sheets']]
if precepts_sheet_id not in existing_sheet_ids:
    batch_reqs.append({
        "addSheet": {
            "properties": {
                "sheetId": precepts_sheet_id,
                "title": "Catalogue Precepts",
                "index": 1,
                "gridProperties": {
                    "rowCount": 805,
                    "columnCount": 56,
                    "frozenRowCount": 1
                }
            }
        }
    })

resp = api_post(f"https://sheets.googleapis.com/v4/spreadsheets/{spreadsheet_id}:batchUpdate", {"requests": batch_reqs})
print("Batch update 1 response:", len(resp.get("replies", [])))

print("\n=== STEP 3: Clear Sheet ID 0 and Populate Enriched Favorites ===")
# Clear existing data on Sheet ID 0 to ensure clean state
api_post(
    f"https://sheets.googleapis.com/v4/spreadsheets/{spreadsheet_id}/values:batchClear",
    {"ranges": ["'Enriched Favorites'!A1:ZZ1000"]}
)
print("Cleared Enriched Favorites sheet range.")

# Load revision_115.xlsx and enriched json reference
with open(enriched_json_path, "r", encoding="utf-8") as f:
    enriched_data = json.load(f)
enriched_map = {item["index"]: item for item in enriched_data}

wb = openpyxl.load_workbook(rev_xlsx_path, data_only=False)
ws = wb.active
headers_rev115 = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
print(f"Read {len(headers_rev115)} headers from revision 115.")

rev115_rows = [headers_rev115]
for r in range(2, ws.max_row + 1):
    row_vals = []
    idx_val = ws.cell(r, 2).value
    try:
        idx = int(idx_val)
    except:
        idx = None
    item = enriched_map.get(idx) if idx else None

    for c in range(1, ws.max_column + 1):
        v = ws.cell(r, c).value
        h = headers_rev115[c - 1]
        
        # Col 1: thumbnail_preview formula
        if c == 1:
            val = str(v) if v is not None else ""
        elif c <= 47 and item and h in item:
            json_v = item.get(h)
            if json_v is None:
                val = ""
            elif isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
                val = str(json_v)
            elif isinstance(v, float) and v.is_integer() and not isinstance(json_v, float):
                val = int(v) if isinstance(json_v, int) else str(json_v)
            elif v is None and json_v is not None:
                val = json_v
            else:
                val = v
        else:
            if v is None:
                val = ""
            elif isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
                val = str(v)
            elif isinstance(v, float) and v.is_integer():
                val = int(v)
            else:
                val = v
        row_vals.append(val)
    rev115_rows.append(row_vals)

print(f"Writing {len(rev115_rows)} rows x {len(rev115_rows[0])} columns to Enriched Favorites...")
val_resp = api_post(
    f"https://sheets.googleapis.com/v4/spreadsheets/{spreadsheet_id}/values:batchUpdate",
    {
        "valueInputOption": "USER_ENTERED",
        "data": [
            {
                "range": f"'Enriched Favorites'!A1:AW{len(rev115_rows)}",
                "majorDimension": "ROWS",
                "values": rev115_rows
            }
        ]
    }
)
print(f"Enriched Favorites populated: {val_resp.get('totalUpdatedCells')} cells.")

print("\n=== STEP 4: Apply Formatting to Enriched Favorites ===")
# Dark Navy Blue (#0D2E66)
navy_color = {"red": 13/255.0, "green": 46/255.0, "blue": 102/255.0}

fmt_reqs_rev115 = [
    # Header format: Dark navy blue, bold white text, centered, middle
    {
        "repeatCell": {
            "range": {
                "sheetId": 0,
                "startRowIndex": 0,
                "endRowIndex": 1,
                "startColumnIndex": 0,
                "endColumnIndex": 49
            },
            "cell": {
                "userEnteredFormat": {
                    "backgroundColor": navy_color,
                    "textFormat": {"foregroundColor": {"red": 1.0, "green": 1.0, "blue": 1.0}, "bold": True, "fontSize": 10},
                    "horizontalAlignment": "CENTER",
                    "verticalAlignment": "MIDDLE"
                }
            },
            "fields": "userEnteredFormat(backgroundColor,textFormat,horizontalAlignment,verticalAlignment)"
        }
    },
    # Set Column A width to 85px
    {
        "updateDimensionProperties": {
            "range": {
                "sheetId": 0,
                "dimension": "COLUMNS",
                "startIndex": 0,
                "endIndex": 1
            },
            "properties": {"pixelSize": 85},
            "fields": "pixelSize"
        }
    },
    # Set row heights for all data rows to 55px (rows 2 to 822)
    {
        "updateDimensionProperties": {
            "range": {
                "sheetId": 0,
                "dimension": "ROWS",
                "startIndex": 1,
                "endIndex": 822
            },
            "properties": {"pixelSize": 55},
            "fields": "pixelSize"
        }
    },
    # Center index column (Col B)
    {
        "repeatCell": {
            "range": {
                "sheetId": 0,
                "startRowIndex": 1,
                "endRowIndex": 822,
                "startColumnIndex": 1,
                "endColumnIndex": 2
            },
            "cell": {
                "userEnteredFormat": {"horizontalAlignment": "CENTER"}
            },
            "fields": "userEnteredFormat(horizontalAlignment)"
        }
    },
    # Auto-resize columns B through AW (col indices 1 through 49)
    {
        "autoResizeDimensions": {
            "dimensions": {
                "sheetId": 0,
                "dimension": "COLUMNS",
                "startIndex": 1,
                "endIndex": 49
            }
        }
    }
]

api_post(f"https://sheets.googleapis.com/v4/spreadsheets/{spreadsheet_id}:batchUpdate", {"requests": fmt_reqs_rev115})
print("Enriched Favorites formatting applied successfully!")

print("\n=== STEP 5: Populate Catalogue Precepts Tab ===")
with open(analyzed_json_path, "r", encoding="utf-8") as f:
    analyzed_data = json.load(f)

analyzed_cols = list(analyzed_data[0].keys())
print(f"Loaded {len(analyzed_data)} analyzed records with {len(analyzed_cols)} columns.")

precepts_rows = [analyzed_cols]
for r in analyzed_data:
    row_vals = [("" if r[k] is None else r[k]) for k in analyzed_cols]
    precepts_rows.append(row_vals)

# Clear Catalogue Precepts if needed
api_post(
    f"https://sheets.googleapis.com/v4/spreadsheets/{spreadsheet_id}/values:batchClear",
    {"ranges": ["'Catalogue Precepts'!A1:ZZ1000"]}
)

# Populate values
val_resp2 = api_post(
    f"https://sheets.googleapis.com/v4/spreadsheets/{spreadsheet_id}/values:batchUpdate",
    {
        "valueInputOption": "USER_ENTERED",
        "data": [
            {
                "range": f"'Catalogue Precepts'!A1:BD{len(precepts_rows)}",
                "majorDimension": "ROWS",
                "values": precepts_rows
            }
        ]
    }
)
print(f"Catalogue Precepts populated: {val_resp2.get('totalUpdatedCells')} cells.")

print("\n=== STEP 6: Apply Professional Formatting to Catalogue Precepts ===")
# Deep slate / navy header (#1E293B)
slate_navy = {"red": 30/255.0, "green": 41/255.0, "blue": 59/255.0}
# Light gray border (#E2E8F0)
border_gray = {"red": 226/255.0, "green": 232/255.0, "blue": 240/255.0}
solid_border = {"style": "SOLID", "color": border_gray}

col_specs = {
    "thumbnail_preview": {"width": 85, "align": "CENTER"},
    "index": {"width": 75, "align": "CENTER"},
    "title": {"width": 240, "align": "LEFT"},
    "original_title": {"width": 200, "align": "LEFT"},
    "creator": {"width": 220, "align": "LEFT"},
    "creator_lifespan": {"width": 110, "align": "CENTER"},
    "creator_nationality": {"width": 120, "align": "CENTER"},
    "date_created": {"width": 90, "align": "CENTER"},
    "location_created": {"width": 160, "align": "LEFT"},
    "medium": {"width": 200, "align": "LEFT"},
    "object_type": {"width": 160, "align": "LEFT"},
    "aspect_ratio": {"width": 85, "align": "CENTER"},
    "aspect_ratio_standard": {"width": 95, "align": "CENTER"},
    "orientation": {"width": 90, "align": "CENTER"},
    "master_width_px": {"width": 85, "align": "CENTER"},
    "master_height_px": {"width": 85, "align": "CENTER"},
    "megapixels": {"width": 85, "align": "CENTER"},
    "focal_anchors_count": {"width": 85, "align": "CENTER"},
    "physical_dimensions_raw": {"width": 160, "align": "LEFT"},
    "parsed_height_cm": {"width": 85, "align": "CENTER"},
    "parsed_width_cm": {"width": 85, "align": "CENTER"},
    "parsed_depth_cm": {"width": 85, "align": "CENTER"},
    "resolution_density_dpi": {"width": 95, "align": "CENTER"},
    "partner_name": {"width": 220, "align": "LEFT"},
    "partner_city_country": {"width": 150, "align": "LEFT"},
    "partner_website": {"width": 160, "align": "LEFT"},
    "partner_lat": {"width": 85, "align": "CENTER"},
    "partner_long": {"width": 85, "align": "CENTER"},
    "curatorial_description": {"width": 360, "align": "LEFT", "wrap": True},
    "provenance": {"width": 240, "align": "LEFT", "wrap": True},
    "credit_line": {"width": 200, "align": "LEFT", "wrap": True},
    "inventory_number": {"width": 110, "align": "CENTER"},
    "external_catalog_url": {"width": 150, "align": "LEFT"},
    "rights": {"width": 180, "align": "LEFT"},
    "art_movements": {"width": 160, "align": "LEFT"},
    "tags_and_topics": {"width": 240, "align": "LEFT", "wrap": True},
    "dominant_color": {"width": 110, "align": "CENTER", "mono": True},
    "wikidata_qid": {"width": 95, "align": "CENTER"},
    "wikipedia_url": {"width": 160, "align": "LEFT"},
    "commons_image_url": {"width": 160, "align": "LEFT"},
    "asset_id": {"width": 120, "align": "CENTER"},
    "link": {"width": 160, "align": "LEFT"},
    "image_url": {"width": 160, "align": "LEFT"},
    "max_preview_url": {"width": 160, "align": "LEFT"},
    "local_image_file": {"width": 200, "align": "LEFT"},
    "gdrive_file_id": {"width": 140, "align": "CENTER"},
    "gdrive_image_url": {"width": 160, "align": "LEFT"},
    "precept_critique": {"width": 480, "align": "LEFT", "wrap": True},
    "design_heuristics": {"width": 520, "align": "LEFT", "wrap": True},
    "extracted_palette": {"width": 220, "align": "LEFT", "mono": True},
    "chromatic_temperature": {"width": 130, "align": "CENTER"},
    "light_profile": {"width": 120, "align": "CENTER"},
    "contrast_level": {"width": 110, "align": "CENTER"},
    "focal_flow": {"width": 360, "align": "LEFT", "wrap": True},
    "framing_density": {"width": 120, "align": "CENTER"},
    "spatial_depth": {"width": 130, "align": "CENTER"}
}

fmt_reqs_precepts = [
    # 1. Header row styling: Deep slate/navy (#1E293B), bold white text, centered, middle
    {
        "repeatCell": {
            "range": {
                "sheetId": precepts_sheet_id,
                "startRowIndex": 0,
                "endRowIndex": 1,
                "startColumnIndex": 0,
                "endColumnIndex": 56
            },
            "cell": {
                "userEnteredFormat": {
                    "backgroundColor": slate_navy,
                    "textFormat": {"foregroundColor": {"red": 1.0, "green": 1.0, "blue": 1.0}, "bold": True, "fontSize": 10},
                    "horizontalAlignment": "CENTER",
                    "verticalAlignment": "MIDDLE"
                }
            },
            "fields": "userEnteredFormat(backgroundColor,textFormat,horizontalAlignment,verticalAlignment)"
        }
    },
    # 2. Header row height: 36px
    {
        "updateDimensionProperties": {
            "range": {
                "sheetId": precepts_sheet_id,
                "dimension": "ROWS",
                "startIndex": 0,
                "endIndex": 1
            },
            "properties": {"pixelSize": 36},
            "fields": "pixelSize"
        }
    },
    # 3. Default vertical alignment TOP for all data rows (rows 1..802)
    {
        "repeatCell": {
            "range": {
                "sheetId": precepts_sheet_id,
                "startRowIndex": 1,
                "endRowIndex": 802,
                "startColumnIndex": 0,
                "endColumnIndex": 56
            },
            "cell": {
                "userEnteredFormat": {
                    "verticalAlignment": "TOP"
                }
            },
            "fields": "userEnteredFormat(verticalAlignment)"
        }
    },
    # 4. Alternating row banding (#FFFFFF / #F8FAFC)
    {
        "addBanding": {
            "bandedRange": {
                "range": {
                    "sheetId": precepts_sheet_id,
                    "startRowIndex": 0,
                    "endRowIndex": 802,
                    "startColumnIndex": 0,
                    "endColumnIndex": 56
                },
                "rowProperties": {
                    "headerColor": slate_navy,
                    "firstBandColor": {"red": 1.0, "green": 1.0, "blue": 1.0},
                    "secondBandColor": {"red": 248/255.0, "green": 250/255.0, "blue": 252/255.0}
                }
            }
        }
    },
    # 5. Light gray cell borders (#E2E8F0) on the entire data grid
    {
        "updateBorders": {
            "range": {
                "sheetId": precepts_sheet_id,
                "startRowIndex": 0,
                "endRowIndex": 802,
                "startColumnIndex": 0,
                "endColumnIndex": 56
            },
            "top": solid_border,
            "bottom": solid_border,
            "left": solid_border,
            "right": solid_border,
            "innerHorizontal": solid_border,
            "innerVertical": solid_border
        }
    }
]

# 6. Add explicit column widths and cell alignments / wrapping for each of the 56 columns
for col_idx, col_name in enumerate(analyzed_cols):
    spec = col_specs[col_name]
    width = spec["width"]
    align = spec["align"]
    is_wrap = spec.get("wrap", False)
    is_mono = spec.get("mono", False)
    
    # Width request
    fmt_reqs_precepts.append({
        "updateDimensionProperties": {
            "range": {
                "sheetId": precepts_sheet_id,
                "dimension": "COLUMNS",
                "startIndex": col_idx,
                "endIndex": col_idx + 1
            },
            "properties": {"pixelSize": width},
            "fields": "pixelSize"
        }
    })
    
    # Cell formatting for column data rows
    cell_format = {
        "horizontalAlignment": align,
        "verticalAlignment": "TOP"
    }
    fields_to_update = ["horizontalAlignment", "verticalAlignment"]
    
    if is_wrap:
        cell_format["wrapStrategy"] = "WRAP"
        fields_to_update.append("wrapStrategy")
        
    if is_mono:
        cell_format["textFormat"] = {"fontFamily": "Roboto Mono", "fontSize": 9}
        fields_to_update.append("textFormat")
        
    fmt_reqs_precepts.append({
        "repeatCell": {
            "range": {
                "sheetId": precepts_sheet_id,
                "startRowIndex": 1,
                "endRowIndex": 802,
                "startColumnIndex": col_idx,
                "endColumnIndex": col_idx + 1
            },
            "cell": {
                "userEnteredFormat": cell_format
            },
            "fields": f"userEnteredFormat({','.join(fields_to_update)})"
        }
    })

print(f"Submitting {len(fmt_reqs_precepts)} formatting requests for Catalogue Precepts...")
api_post(f"https://sheets.googleapis.com/v4/spreadsheets/{spreadsheet_id}:batchUpdate", {"requests": fmt_reqs_precepts})
print("Catalogue Precepts formatted successfully!")

print("\n=== COMPLETE ===")
